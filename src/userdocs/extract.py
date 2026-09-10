"""Text extraction for uploaded documents. See spec/v3/SPEC.md §5.1."""

import io
from dataclasses import dataclass

from docx import Document as DocxDocument
from pypdf import PdfReader

from userdocs.errors import CorruptDocumentError, EmptyDocumentError, UnsupportedFormatError

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


@dataclass
class ExtractedDocument:
    text: str
    pages: list[str] | None


def _extension_of(filename: str) -> str:
    idx = filename.rfind(".")
    return filename[idx:].lower() if idx != -1 else ""


def extract_text(filename: str, raw_bytes: bytes) -> ExtractedDocument:
    ext = _extension_of(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(f"Unsupported file extension: {ext!r}")

    if ext in (".txt", ".md"):
        text = raw_bytes.decode("utf-8", errors="replace")
        result = ExtractedDocument(text=text, pages=None)
    elif ext == ".docx":
        try:
            doc = DocxDocument(io.BytesIO(raw_bytes))
        except Exception as exc:
            # Deliberately broad: python-docx raises PackageNotFoundError for a
            # non-zip file, but a truncated or subtly-malformed .docx surfaces
            # as KeyError, BadZipFile, or an lxml parse error depending on which
            # part is damaged. The caller only needs "this file is broken" — so
            # every parse-time failure maps to one exception type rather than
            # enumerating a list that a new python-docx version can invalidate.
            raise CorruptDocumentError(f"{filename} could not be parsed as .docx") from exc
        text = "\n".join(p.text for p in doc.paragraphs)
        result = ExtractedDocument(text=text, pages=None)
    else:  # ".pdf"
        try:
            reader = PdfReader(io.BytesIO(raw_bytes))
            # `reader.pages` is lazy — page text is only extracted (and only
            # fails) when iterated, so the iteration has to be INSIDE the try.
            pages = [(page.extract_text() or "") for page in reader.pages]
        except Exception as exc:
            # Broad for the same reason as .docx above: pypdf raises
            # PdfReadError for a bad trailer but other exception types for
            # damaged xref tables, bad encodings, or encrypted files.
            raise CorruptDocumentError(f"{filename} could not be parsed as .pdf") from exc
        result = ExtractedDocument(text="\n".join(pages), pages=pages)

    if not result.text.strip():
        raise EmptyDocumentError(f"{filename} contains no extractable text")
    return result
